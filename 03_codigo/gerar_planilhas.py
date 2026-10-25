# -*- coding: utf-8 -*-
"""Gera a planilha editável com fórmulas: 04_planilhas/PT_calculos_Vivara_241003384.xlsx
Executar depois de gerar_base.py, testes_auditoria.py e calculos_pts.py."""
import json, os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import dados_publicos as D
import parametros as P
from distorcoes import PUBLICAS

BASE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(BASE, "..", "02_base_simulada", "resultados")
OUT = os.path.join(BASE, "..", "04_planilhas")
os.makedirs(OUT, exist_ok=True)
calc = json.load(open(os.path.join(RES, "resultados_calculos.json")))
tst = json.load(open(os.path.join(RES, "resultados_testes.json")))

wb = Workbook()
H = PatternFill("solid", fgColor="1F3864"); HF = Font(color="FFFFFF", bold=True)
INP = PatternFill("solid", fgColor="FFF2CC"); SIM = PatternFill("solid", fgColor="FCE4D6")
B = Font(bold=True); thin = Side(style="thin", color="BFBFBF")
NUM = '#,##0;[Red]-#,##0'; NUM2 = '#,##0.00;[Red]-#,##0.00'; PCT = '0.00%'


def header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c); x.fill = H; x.font = HF
        x.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, w):
    for j, v in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(j)].width = v


def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")
    ws["A3"] = "Preparadora: Rafaela de Sousa Barbosa - matrícula 241003384 | Data: 22/10/2026 | Conteúdo produzido com IA (ver README) | Revisão: ____________ (campo não preenchido)"
    ws["A3"].font = Font(size=8, color="7F7F7F")


# ---------------------------------------------------------------- LEIA-ME
ws = wb.active; ws.title = "LEIA-ME"
linhas = ["Vivara Participações S.A. - Dossiê didático de auditoria (exercício 2025) - planilha de cálculos",
          "UnB - Auditoria 1 (CCA0024) - Matrícula 241003384",
          "Valores em R$ mil, exceto quando indicado (R$ nas abas da base simulada).",
          "Células amarelas = dados de entrada públicos (com fonte). Células laranja = dados SINTÉTICOS (base simulada).",
          "Demais células numéricas contêm fórmulas e podem ser alteradas para análise de sensibilidade.",
          "Classificação dos procedimentos: Realizado com evidência pública | Planejado por falta de acesso | Simulado para fins didáticos.",
          "Nenhum dado sintético representa registro real da companhia.",
          "Abas: PT320 | PT500 | PT510 | PT520 | PT530_MUS | PT530_Amostra | PT530_Testes100 | PT450 | PT540_Arrend | PT540_ProvEst | PT570 | PT600 | PT505"]
for i, l in enumerate(linhas, 1):
    ws.cell(row=i, column=1, value=l)
ws["A1"].font = Font(bold=True, size=13); widths(ws, [130])

# ---------------------------------------------------------------- PT320
ws = wb.create_sheet("PT320")
title(ws, "PT 320 - Materialidade", "NBC TA 320 (R1) itens 10-14 e A3-A13 | Realizado com evidência pública")
header(ws, 5, ["Referencial", "Base (R$ mil)", "% aplicado", "Materialidade (R$ mil)", "Fonte"])
fontes = {"LAIR 2025": "DF25 p.20", "LAIR médio 2023-2025": "DF25 p.20; DFP24 p.19", "Receita líquida 2025": "DF25 p.20",
          "Receita líquida 2025 (0,5%)": "DF25 p.20", "Ativo total 2025": "DF25 p.19", "Patrimônio líquido 2025": "DF25 p.19",
          "Lucro bruto 2025": "DF25 p.20"}
r = 6
for a in calc["PT320"]["alternativas"]:
    ws.cell(row=r, column=1, value=a["referencial"])
    c = ws.cell(row=r, column=2, value=a["base"]); c.fill = INP; c.number_format = NUM
    c = ws.cell(row=r, column=3, value=a["percentual"]); c.number_format = PCT; c.fill = INP
    c = ws.cell(row=r, column=4, value=f"=B{r}*C{r}"); c.number_format = NUM
    ws.cell(row=r, column=5, value=fontes[a["referencial"]]); r += 1
r += 1
ws.cell(row=r, column=1, value="Materialidade global calculada (5% LAIR 2025)").font = B
ws.cell(row=r, column=4, value="=D6").number_format = NUM; rmg = r; r += 1
ws.cell(row=r, column=1, value="Materialidade global adotada (arredondada para baixo)").font = B
c = ws.cell(row=r, column=4, value=P.MG); c.number_format = NUM; c.fill = INP; rmga = r; r += 1
ws.cell(row=r, column=1, value="% materialidade para execução (auditoria inicial)")
c = ws.cell(row=r, column=4, value=P.PERC_ME); c.number_format = PCT; c.fill = INP; rpme = r; r += 1
ws.cell(row=r, column=1, value="Materialidade para execução").font = B
ws.cell(row=r, column=4, value=f"=D{rmga}*D{rpme}").number_format = NUM; r += 1
ws.cell(row=r, column=1, value="% limite claramente trivial")
c = ws.cell(row=r, column=4, value=P.PERC_TRIVIAL); c.number_format = PCT; c.fill = INP; rpt = r; r += 1
ws.cell(row=r, column=1, value="Limite para distorções claramente triviais").font = B
ws.cell(row=r, column=4, value=f"=D{rmga}*D{rpt}").number_format = NUM; r += 2
ws.cell(row=r, column=1, value="Materialidades específicas").font = B; r += 1
for e in calc["PT320"]["especificas"]:
    ws.cell(row=r, column=1, value=e["item"]); ws.cell(row=r, column=4, value=e["materialidade"]).number_format = NUM
    ws.cell(row=r, column=5, value=e["base"]); r += 1
r += 1
ws.cell(row=r, column=1, value="Indicadores da MG adotada").font = B; r += 1
for lab, base in [("MG / receita líquida", D.DRE["2025"]["receita_liquida"]), ("MG / ativo total", D.BP["2025"]["ativo_total"]),
                  ("MG / patrimônio líquido", D.BP["2025"]["patrimonio_liquido"]), ("MG / lucro líquido", D.DRE["2025"]["lucro_liquido"])]:
    ws.cell(row=r, column=1, value=lab); c = ws.cell(row=r, column=2, value=base); c.fill = INP; c.number_format = NUM
    ws.cell(row=r, column=4, value=f"=D{rmga}/B{r}").number_format = PCT; r += 1
widths(ws, [52, 16, 12, 22, 60])

# ---------------------------------------------------------------- PT500
ws = wb.create_sheet("PT500")
title(ws, "PT 500 - Conciliações e verificações de consistência das demonstrações (R$ mil)",
      "NBC TA 500 (R1) itens 6-9, A26-A33 | Realizado com evidência pública")
header(ws, 5, ["Verificação", "Calculado", "Divulgado", "Diferença", "Status", "Fonte"])
r = 6
for c in calc["PT500"]:
    ws.cell(row=r, column=1, value=c["verificacao"])
    x = ws.cell(row=r, column=2, value=c["calculado"]); x.number_format = NUM2
    x = ws.cell(row=r, column=3, value=c["divulgado"]); x.number_format = NUM2; x.fill = INP
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = NUM2
    ws.cell(row=r, column=5, value=c["status"]); ws.cell(row=r, column=6, value=c["fonte"])
    if c["status"] != "OK":
        ws.cell(row=r, column=5).font = Font(color="C00000", bold=True)
    r += 1
r += 1
ws.cell(row=r, column=1, value="Conciliação da base simulada com o saldo contábil (R$) - SIMULADO").font = B; r += 1
header(ws, r, ["Item", "Base simulada (R$)", "Contábil (R$)", "Diferença", "", "Fonte contábil"]); r += 1
for c in tst["conciliacao"]:
    ws.cell(row=r, column=1, value=c["item"])
    x = ws.cell(row=r, column=2, value=c["base_simulada_RS"]); x.number_format = NUM2; x.fill = SIM
    x = ws.cell(row=r, column=3, value=c["contabil_RS"]); x.number_format = NUM2; x.fill = INP
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = NUM2
    ws.cell(row=r, column=6, value=c["fonte_contabil"]); r += 1
widths(ws, [78, 18, 18, 14, 14, 40])

# ---------------------------------------------------------------- PT510
ws = wb.create_sheet("PT510")
title(ws, "PT 510 - Saldos iniciais: DFP 2024 (formulário CVM) x comparativos das DF 2025 (R$ mil)",
      "NBC TA 510 (R1) itens 5-9 | Realizado com evidência pública")
header(ws, 5, ["Linha", "DFP 2024 (R$ mil)", "Comparativo DF 2025", "Diferença", "Status"])
r = 6
for c in calc["PT510"]:
    ws.cell(row=r, column=1, value=c["linha"])
    x = ws.cell(row=r, column=2, value=c["DFP2024_RS_mil"]); x.number_format = NUM2; x.fill = INP
    x = ws.cell(row=r, column=3, value=c["comparativo_DF2025_RS_mil"]); x.number_format = NUM; x.fill = INP
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = NUM2
    ws.cell(row=r, column=5, value=f'=IF(ABS(D{r})<=1.5,"OK (arredondamento)","DIVERGÊNCIA")'); r += 1
r += 1
ws.cell(row=r, column=1, value="Retificação 2023 (nota 3 da DFP 2024) x formulário DFP 2024 - lucro líquido 2023").font = B; r += 1
x = ws.cell(row=r, column=1, value="Lucro 2023 no formulário DFP 2024 (DFP24 p.19)"); c = ws.cell(row=r, column=2, value=calc["PT510_retificacao_2023"]["lucro_2023_formulario_DFP24"]); c.fill = INP; c.number_format = NUM2; r += 1
ws.cell(row=r, column=1, value="Lucro 2023 retificado (DFP24 nota 3, p.60)"); c = ws.cell(row=r, column=2, value=calc["PT510_retificacao_2023"]["lucro_2023_retificado_nota3"]); c.fill = INP; c.number_format = NUM2; r += 1
ws.cell(row=r, column=1, value="Diferença"); ws.cell(row=r, column=2, value=f"=B{r-1}-B{r-2}").number_format = NUM2
widths(ws, [62, 18, 20, 14, 22])

# ---------------------------------------------------------------- PT520
ws = wb.create_sheet("PT520")
title(ws, "PT 520 - Procedimentos analíticos substantivos (R$ mil)", "NBC TA 520 itens 5-7, A4-A16 | Realizado com evidência pública")
ws["A5"] = "Limite de investigação (= materialidade para execução)"; ws["B5"] = P.ME; ws["B5"].fill = INP; ws["B5"].number_format = NUM
R25, R24 = D.DRE["2025"], D.DRE["2024"]
inputs = [("Receita líquida 2025", R25["receita_liquida"], "DF25 p.20"), ("Receita líquida 2024", R24["receita_liquida"], "DF25 p.20"),
          ("CMV 2025", -R25["cmv"], "DF25 p.20"), ("CMV 2024", -R24["cmv"], "DF25 p.20"),
          ("Despesas com vendas 2025", -R25["desp_vendas"], "DF25 p.20"), ("Despesas com vendas 2024", -R24["desp_vendas"], "DF25 p.20"),
          ("G&A 2025", -R25["desp_gerais_adm"], "DF25 p.20"), ("G&A 2024", -R24["desp_gerais_adm"], "DF25 p.20"),
          ("Receita lojas físicas 2024", D.RECEITA["canais_2024"]["Lojas"], "DF25 p.55"),
          ("Receita lojas físicas 2025", D.RECEITA["canais_2025"]["Lojas"], "DF25 p.55"),
          ("SSS 2025", D.OUTROS["sss_2025"], "DF25 p.3"), ("Aberturas líquidas 2025", 41, "DF25 p.4"),
          ("Receita média anual loja Life madura", 5600, "DF25 p.7"), ("Fração do ano (lojas novas)", 0.5, "premissa"),
          ("GGF alocado ao estoque 2024 (RA)", 43900, "DF25 p.8"), ("GGF alocado ao estoque 2025 (RA)", 24400, "DF25 p.8"),
          ("Passivo de arrendamento 2024", D.ARREND["passivo_2024"], "DF25 p.59"), ("Passivo de arrendamento 2025", D.ARREND["passivo_2025"], "DF25 p.59"),
          ("Taxa média de desconto 2025", D.ARREND["taxa_media"], "DF25 p.58"), ("Encargos de arrendamento 2025", D.ARREND["encargos_2025"], "DF25 p.57"),
          ("Debêntures - principal", 300000, "DF25 p.40"), ("CDI proxy (cenário I nota 29.7)", 0.1432, "DF25 p.66"),
          ("Spread debêntures", 0.007, "DF25 p.40"), ("Dias 05/09 a 31/12/2025", 117, "cálculo"),
          ("Juros debêntures incorridos 2025", 14401, "DF25 p.41")]
header(ws, 7, ["Dado de entrada", "Valor", "Fonte"])
pos = {}
for i, (n, v, f) in enumerate(inputs, start=8):
    ws.cell(row=i, column=1, value=n); c = ws.cell(row=i, column=2, value=v); c.fill = INP
    c.number_format = PCT if isinstance(v, float) and v < 1 else NUM
    ws.cell(row=i, column=3, value=f); pos[n] = f"B{i}"
r = 8 + len(inputs) + 1
header(ws, r, ["Conta", "Expectativa (fórmula)", "Real", "Diferença", "Dif. %", "Investigar?", "Modelo"]); r += 1
p = pos
mods = [("Receita bruta - lojas físicas", f"={p['Receita lojas físicas 2024']}*(1+{p['SSS 2025']})+{p['Aberturas líquidas 2025']}*{p['Receita média anual loja Life madura']}*{p['Fração do ano (lojas novas)']}",
         f"={p['Receita lojas físicas 2025']}", "Lojas 2024 x (1+SSS) + lojas novas x receita média x fração do ano"),
        ("CMV", f"={p['Receita líquida 2025']}*{p['CMV 2024']}/{p['Receita líquida 2024']}", f"={p['CMV 2025']}", "RL 2025 x (CMV/RL 2024)"),
        ("CMV - ajustado pelo efeito GGF", f"={p['Receita líquida 2025']}*({p['CMV 2024']}+{p['GGF alocado ao estoque 2024 (RA)']})/{p['Receita líquida 2024']}-{p['GGF alocado ao estoque 2025 (RA)']}",
         f"={p['CMV 2025']}", "Investigação: margem 2024 sem alocação de GGF; deduz GGF alocado em 2025"),
        ("Despesas com vendas", f"={p['Receita líquida 2025']}*{p['Despesas com vendas 2024']}/{p['Receita líquida 2024']}", f"={p['Despesas com vendas 2025']}", "RL 2025 x (DV/RL 2024)"),
        ("Despesas gerais e administrativas", f"={p['Receita líquida 2025']}*{p['G&A 2024']}/{p['Receita líquida 2024']}", f"={p['G&A 2025']}", "RL 2025 x (G&A/RL 2024)"),
        ("Encargos financeiros de arrendamento", f"=({p['Passivo de arrendamento 2024']}+{p['Passivo de arrendamento 2025']})/2*{p['Taxa média de desconto 2025']}",
         f"={p['Encargos de arrendamento 2025']}", "Passivo médio x taxa média"),
        ("Juros das debêntures", f"={p['Debêntures - principal']}*(((1+{p['CDI proxy (cenário I nota 29.7)']})*(1+{p['Spread debêntures']}))^({p['Dias 05/09 a 31/12/2025']}/365)-1)",
         f"={p['Juros debêntures incorridos 2025']}", "Principal x fator DI+spread no período")]
for n, e, re_, m in mods:
    ws.cell(row=r, column=1, value=n); ws.cell(row=r, column=2, value=e).number_format = NUM
    ws.cell(row=r, column=3, value=re_).number_format = NUM
    ws.cell(row=r, column=4, value=f"=C{r}-B{r}").number_format = NUM
    ws.cell(row=r, column=5, value=f"=D{r}/B{r}").number_format = PCT
    ws.cell(row=r, column=6, value=f'=IF(ABS(D{r})>$B$5,"SIM","NÃO")'); ws.cell(row=r, column=7, value=m); r += 1
r += 1
ws.cell(row=r, column=1, value="Índices (calculados em Python - resultados_calculos.json)").font = B; r += 1
for k, v in calc["PT520_indices"].items():
    ws.cell(row=r, column=1, value=k); ws.cell(row=r, column=2, value=json.dumps(v, ensure_ascii=False)); r += 1
widths(ws, [44, 22, 16, 14, 10, 12, 70])

# ---------------------------------------------------------------- PT530 MUS
ws = wb.create_sheet("PT530_MUS")
mus = tst["mus"]
title(ws, "PT 530 - Amostragem por unidade monetária (MUS) - BASE SIMULADA (R$)", "NBC TA 530 itens 6-15, A10-A23 | SIMULADO PARA FINS DIDÁTICOS")
rows = [("Valor contábil da população (itens > 0) - R$", mus["valor_populacao_RS"], SIM, NUM2),
        ("Nº de itens da população", mus["populacao_itens"], SIM, NUM),
        ("Distorção tolerável (= ME) - R$", mus["distorcao_toleravel_RS"], INP, NUM2),
        ("% distorção esperada", P.PERC_DISTORCAO_ESPERADA, INP, PCT),
        ("Distorção esperada - R$", "=B8*B9", None, NUM2),
        ("Fator de confiabilidade (95%, 0 erros)", P.FATOR_CONFIABILIDADE, INP, '0.00'),
        ("Fator de expansão", P.FATOR_EXPANSAO, INP, '0.00'),
        ("Tamanho da amostra (n)", "=ROUNDUP(B6*B11/(B8-B10*B12),0)", None, NUM),
        ("Intervalo amostral (J) - R$", "=B6/B13", None, NUM2),
        ("Início aleatório - R$ (semente 241003384)", mus["inicio_aleatorio_RS"], SIM, NUM2),
        ("Itens com distorção na amostra", mus["itens_com_distorcao_total"], SIM, NUM),
        ("  - já identificados nos testes de 100%", mus["itens_com_distorcao_ja_identificada_100pct"], SIM, NUM),
        ("  - projetáveis", mus["itens_com_distorcao_projetada"], SIM, NUM),
        ("Precisão básica (J x 3,0) - R$", "=B14*B11", None, NUM2),
        ("Distorção projetada - R$", mus["distorcao_projetada_RS"], SIM, NUM2),
        ("Provisão incremental - R$", mus["provisao_incremental_RS"], SIM, NUM2),
        ("Limite superior de distorção (UML) - R$", "=B19+B20+B21", None, NUM2),
        ("Conclusão", '=IF(B22<=B8,"UML <= tolerável: aceitar","UML > tolerável: ampliar/ajustar")', None, None)]
for i, (lab, val, fill, fmt) in enumerate(rows, start=6):
    ws.cell(row=i, column=1, value=lab); c = ws.cell(row=i, column=2, value=val)
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
widths(ws, [52, 24])

ws = wb.create_sheet("PT530_Amostra")
title(ws, "PT 530 - Itens selecionados (MUS) e resultado da contagem/vouching SIMULADOS (R$)", "Dados SINTÉTICOS")
am = pd.read_csv(os.path.join(RES, "PT530_amostra_MUS_selecionada.csv"), sep=";", decimal=",")
cols = ["id_registro", "cod_local", "tipo_local", "classe_estoque", "categoria", "cod_sku", "quantidade", "custo_unitario",
        "valor_liquido", "ponto_selecao_RS", "quantidade_contada_simulada", "custo_documental_simulado", "valor_auditado_RS",
        "distorcao_RS", "taint", "ja_identificado_teste_100pct"]
header(ws, 5, cols)
for i, rr in enumerate(am[cols].itertuples(index=False), start=6):
    for j, v in enumerate(rr, 1):
        c = ws.cell(row=i, column=j, value=(bool(v) if isinstance(v, (bool,)) else v)); c.fill = SIM
widths(ws, [18, 10, 20, 22, 18, 14, 10, 12, 14, 14, 12, 12, 14, 12, 10, 12])

ws = wb.create_sheet("PT530_Testes100")
title(ws, "PT 530/501 - Testes em 100% da base simulada e investigação dos itens sinalizados (R$)", "SIMULADO PARA FINS DIDÁTICOS")
header(ws, 5, ["Teste", "Itens sinalizados", "Confirmados", "Distorção confirmada (R$)"])
r = 6
for t, v in sorted(tst["testes_100pct"].items()):
    ws.cell(row=r, column=1, value=t); ws.cell(row=r, column=2, value=v["sinalizados"]); ws.cell(row=r, column=3, value=v["confirmados"])
    ws.cell(row=r, column=4, value=v["distorcao_confirmada_RS"]).number_format = NUM2; r += 1
r += 1
ws.cell(row=r, column=1, value="Comparação erros inseridos x identificados").font = B; r += 1
header(ws, r, ["Tipo de erro", "Inseridos", "Detectados", "Taxa", "Efeito inserido (R$)", "Distorção quantificada (R$)", "Detectados 100%", "Detectados amostra"]); r += 1
for c in tst["comparacao"]:
    ws.cell(row=r, column=1, value=c["tipo_erro"]); ws.cell(row=r, column=2, value=c["inseridos"]); ws.cell(row=r, column=3, value=c["detectados"])
    ws.cell(row=r, column=4, value=f"=C{r}/B{r}").number_format = PCT
    ws.cell(row=r, column=5, value=c["efeito_inserido_RS"]).number_format = NUM2
    ws.cell(row=r, column=6, value=c["distorcao_quantificada_RS"]).number_format = NUM2
    ws.cell(row=r, column=7, value=c["detectados_100pct"]); ws.cell(row=r, column=8, value=c["detectados_amostra"]); r += 1
widths(ws, [44, 14, 14, 10, 20, 24, 14, 16])

# ---------------------------------------------------------------- PT450
ws = wb.create_sheet("PT450")
title(ws, "PT 450 - Mapa de distorções (R$ mil)", "NBC TA 450 (R1) itens 5-12 | Público = evidência pública | Simulado = base sintética")
ws["A5"] = "MG"; ws["B5"] = P.MG; ws["A6"] = "ME"; ws["B6"] = P.ME; ws["A7"] = "Limite trivial"; ws["B7"] = P.LIMITE_TRIVIAL
for c in ("B5", "B6", "B7"):
    ws[c].number_format = NUM; ws[c].fill = INP
header(ws, 9, ["Código", "Origem", "Descrição", "Fonte", "Natureza", "Valor (R$ mil)", "Afeta DF 2025?", "Tratamento"])
r = 10
for d in PUBLICAS:
    vals = [d[0], "Público", d[1], d[2], d[3], d[4], d[5], d[6]]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=6).number_format = NUM; r += 1
rsim = r
for v in tst["distorcao_por_teste_principal"]:
    t = v["teste_principal"]
    vals = [f"S-{t[:2]}", "Simulado", f"Distorção factual - itens cujo teste principal foi {t} ({v['itens']} itens; base sintética)",
            "PT530_Testes100", "Factual (simulada)", round(v["distorcao_RS"] / 1000, 1), "Simulação", "Não corrigida (cenário didático) - acumulada"]
    for j, x in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=x)
    ws.cell(row=r, column=6).number_format = '#,##0.0'; r += 1
vals = ["S-PJ", "Simulado", "Distorção projetada - MUS (erros de contagem não detectáveis por análise de dados)", "PT530_MUS",
        "Projetada (simulada)", round(mus["distorcao_projetada_RS"] / 1000, 1), "Simulação", "Acumulada"]
for j, x in enumerate(vals, 1):
    ws.cell(row=r, column=j, value=x)
ws.cell(row=r, column=6).number_format = '#,##0.0'; rend = r; r += 2
ws.cell(row=r, column=3, value="Total de distorções simuladas não corrigidas (factuais + projetadas)").font = B
ws.cell(row=r, column=6, value=f"=SUM(F{rsim}:F{rend})").number_format = '#,##0.0'; rt = r; r += 1
ws.cell(row=r, column=3, value="% da materialidade global"); ws.cell(row=r, column=6, value=f"=F{rt}/B5").number_format = PCT; r += 1
ws.cell(row=r, column=3, value="% da materialidade para execução"); ws.cell(row=r, column=6, value=f"=F{rt}/B6").number_format = PCT; r += 1
ws.cell(row=r, column=3, value="Distorções factuais + UML da amostra (risco de amostragem incluído)")
ws.cell(row=r, column=6, value=round((tst["distorcao_conhecida_testes_100pct_RS"] + mus["limite_superior_distorcao_RS"]) / 1000, 1)).number_format = '#,##0.0'; r += 1
ws.cell(row=r, column=3, value="Efeito total dos erros INSERIDOS (gabarito, R$ mil)")
ws.cell(row=r, column=6, value=round(tst["efeito_total_inserido_RS"] / 1000, 1)).number_format = '#,##0.0'
widths(ws, [8, 10, 70, 26, 26, 14, 12, 50])

# ---------------------------------------------------------------- PT540 Arrendamento
ws = wb.create_sheet("PT540_Arrend")
title(ws, "PT 540 - Sensibilidade do passivo de arrendamento à taxa de desconto (R$ mil)",
      "NBC TA 540 (R2) itens 13, 18, 25-26 | Realizado com evidência pública (modelo simplificado)")
fl = D.ARREND["fluxos_nao_descontados_2025"]
cfs = [(0.5, fl["ate_1"], "Até 1 ano"), (1.5, fl["1_2"], "1 a 2 anos"), (2.5, fl["2_3"], "2 a 3 anos"), (3.5, fl["3_5"] / 2, "3 a 5 anos (1/2)"),
       (4.5, fl["3_5"] / 2, "3 a 5 anos (2/2)")] + [(5.5 + i, fl["mais_5"] / 5, f"> 5 anos ({i+1}/5)") for i in range(5)]
ws["A5"] = "Taxas para teste"; rates = [0.1092, 0.1192, 0.1292, 0.1392, 0.1492]
header(ws, 6, ["Faixa", "Tempo (anos)", "Fluxo não descontado"] + [f"VP @ {x:.2%}" for x in rates])
for j, x in enumerate(rates):
    c = ws.cell(row=5, column=4 + j, value=x); c.number_format = PCT; c.fill = INP
r = 7
for t, v, lab in cfs:
    ws.cell(row=r, column=1, value=lab); ws.cell(row=r, column=2, value=t)
    c = ws.cell(row=r, column=3, value=v); c.fill = INP; c.number_format = NUM
    for j in range(len(rates)):
        col = get_column_letter(4 + j)
        ws.cell(row=r, column=4 + j, value=f"=$C{r}/(1+{col}$5)^$B{r}").number_format = NUM
    r += 1
ws.cell(row=r, column=1, value="Total").font = B
ws.cell(row=r, column=3, value=f"=SUM(C7:C{r-1})").number_format = NUM
for j in range(len(rates)):
    col = get_column_letter(4 + j)
    ws.cell(row=r, column=4 + j, value=f"=SUM({col}7:{col}{r-1})").number_format = NUM
rt = r; r += 1
ws.cell(row=r, column=1, value="Passivo divulgado (nota 28)"); c = ws.cell(row=r, column=3, value=D.ARREND["passivo_2025"]); c.fill = INP; c.number_format = NUM; r += 1
ws.cell(row=r, column=1, value="Variação vs taxa divulgada (12,92%)")
for j in range(len(rates)):
    col = get_column_letter(4 + j)
    ws.cell(row=r, column=4 + j, value=f"={col}{rt}-$F${rt}").number_format = NUM
r += 1
ws.cell(row=r, column=1, value="Variação / materialidade para execução")
for j in range(len(rates)):
    col = get_column_letter(4 + j)
    ws.cell(row=r, column=4 + j, value=f"={col}{r-1}/{P.ME}").number_format = PCT
widths(ws, [30, 12, 20, 16, 16, 16, 16, 16])

# ---------------------------------------------------------------- PT540 Provisão de estoques
ws = wb.create_sheet("PT540_ProvEst")
title(ws, "PT 540 - Provisão para perdas de estoques: revisão retrospectiva e cenários (R$ mil)", "Realizado com evidência pública")
pe = calc["PT540_provisao_estoque"]
rows = [("Saldo da provisão 31/12/2025 (nota 7.3)", pe["saldo_31_12_2025"], "DF25 p.31"),
        ("Reversões no 1S26 (ITR 2T26 nota 5.2)", pe["reversoes_1S26"], "ITR2 p.14"),
        ("Complementos no 1S26", pe["complementos_1S26"], "ITR2 p.14"),
        ("Saldo 30/06/2026", pe["saldo_30_06_2026"], "ITR2 p.14"),
        ("Estoque líquido 31/12/2025", D.BP["2025"]["estoques"], "DF25 p.19")]
for i, (a, v, f) in enumerate(rows, start=5):
    ws.cell(row=i, column=1, value=a); c = ws.cell(row=i, column=2, value=v); c.fill = INP; c.number_format = NUM; ws.cell(row=i, column=3, value=f)
ws["A11"] = "% revertido no 1S26"; ws["B11"] = "=B6/B5"; ws["B11"].number_format = PCT
ws["A12"] = "Estoque bruto"; ws["B12"] = "=B9+B5"; ws["B12"].number_format = NUM
ws["A13"] = "Provisão / estoque bruto"; ws["B13"] = "=B5/B12"; ws["B13"].number_format = '0.000%'
header(ws, 15, ["Cenário: provisão como % do estoque bruto", "%", "Provisão", "Efeito no resultado (antes de IR)"])
for i, pp in enumerate([0.005, 0.01, 0.02], start=16):
    ws.cell(row=i, column=1, value=f"Cenário {i-15}"); c = ws.cell(row=i, column=2, value=pp); c.number_format = PCT; c.fill = INP
    ws.cell(row=i, column=3, value=f"=$B$12*B{i}").number_format = NUM
    ws.cell(row=i, column=4, value=f"=-(C{i}-$B$5)").number_format = NUM
widths(ws, [46, 14, 16, 30])

# ---------------------------------------------------------------- PT570
ws = wb.create_sheet("PT570")
title(ws, "PT 570 - Indicadores de liquidez e endividamento (R$ mil)", "NBC TA 570 itens 10-16, A3 | Realizado com evidência pública")
B25, B24 = D.BP["2025"], D.BP["2024"]
header(ws, 5, ["Item", "31/12/2025", "31/12/2024", "Fonte"])
items = [("Ativo circulante", B25["ativo_circulante"], B24["ativo_circulante"]), ("Estoques", B25["estoques"], B24["estoques"]),
         ("Caixa e equivalentes + TVM", B25["caixa_equivalentes"] + B25["tvm"], B24["caixa_equivalentes"] + B24["tvm"]),
         ("Passivo circulante", B25["passivo_circulante"], B24["passivo_circulante"]),
         ("Passivo não circulante", B25["passivo_nao_circulante"], B24["passivo_nao_circulante"]),
         ("Ativo total", B25["ativo_total"], B24["ativo_total"]), ("Patrimônio líquido", B25["patrimonio_liquido"], B24["patrimonio_liquido"]),
         ("Dívida bruta líquida de derivativos (nota 15.1)", D.EMPREST["2025"]["total_liq_deriv"], D.EMPREST["2024"]["total_liq_deriv"]),
         ("Lucro antes do resultado financeiro", D.DRE["2025"]["lucro_antes_res_fin"], D.DRE["2024"]["lucro_antes_res_fin"]),
         ("Depreciação e amortização (DFC)", D.DFC["2025"]["depreciacao_amortizacao"], D.DFC["2024"]["depreciacao_amortizacao"]),
         ("Caixa líquido das atividades operacionais", D.DFC["2025"]["caixa_operacional_liquido"], D.DFC["2024"]["caixa_operacional_liquido"])]
for i, (a, x, y) in enumerate(items, start=6):
    ws.cell(row=i, column=1, value=a)
    for j, v in enumerate([x, y], start=2):
        c = ws.cell(row=i, column=j, value=v); c.fill = INP; c.number_format = NUM
    ws.cell(row=i, column=4, value="DF25 p.19/20/23; nota 15.1 p.40")
r = 6 + len(items) + 1
form = [("Liquidez corrente", "={c}6/{c}9", '0.00'), ("Liquidez seca", "=({c}6-{c}7)/{c}9", '0.00'),
        ("Liquidez imediata", "={c}8/{c}9", '0.00'), ("Endividamento geral (PC+PNC)/AT", "=({c}9+{c}10)/{c}11", PCT),
        ("Dívida líquida", "={c}13-{c}8", NUM), ("LAJIDA calculado", "={c}14+{c}15", NUM),
        ("Dívida líquida / LAJIDA", "=({c}13-{c}8)/({c}14+{c}15)", '0.00')]
for a, fml, fmt in form:
    ws.cell(row=r, column=1, value=a).font = B
    for col in ("B", "C"):
        ws[f"{col}{r}"] = fml.format(c=col); ws[f"{col}{r}"].number_format = fmt
    r += 1
widths(ws, [48, 16, 16, 34])

# ---------------------------------------------------------------- PT600
ws = wb.create_sheet("PT600")
title(ws, "PT 600 - Componentes do grupo e materialidade de componentes (R$ mil)", "NBC TA 600 (R2) | Realizado com evidência pública (nota 9.2)")
header(ws, 5, ["Componente", "Ativo total", "% soma ativos", "Lucro/prejuízo", "% |lucro|", "ME do componente", "Classificação"])
cl = {"Tellerina": "Significativo (varejo, estoques nas lojas)", "Conipa": "Significativo (fábrica, matérias-primas, incentivos)",
      "Tellerina Panama": "Não significativo - procedimentos analíticos no nível do grupo"}
r = 6
for c in calc["PT600"]["componentes"]:
    ws.cell(row=r, column=1, value=c["componente"]); x = ws.cell(row=r, column=2, value=c["ativo_total"]); x.fill = INP; x.number_format = NUM
    ws.cell(row=r, column=3, value=f"=B{r}/SUM($B$6:$B$8)").number_format = PCT
    x = ws.cell(row=r, column=4, value=c["lucro_prejuizo"]); x.fill = INP; x.number_format = NUM
    ws.cell(row=r, column=5, value=f"=ABS(D{r})/SUMPRODUCT(ABS($D$6:$D$8))").number_format = PCT
    ws.cell(row=r, column=6, value=calc["PT600"]["ME_componentes"][c["componente"]]).number_format = NUM
    ws.cell(row=r, column=7, value=cl[c["componente"]]); r += 1
ws.cell(row=r, column=1, value="Controladora (holding)"); ws.cell(row=r, column=6, value=6000).number_format = NUM
ws.cell(row=r, column=7, value="Auditada diretamente pela equipe do grupo"); r += 2
ws.cell(row=r, column=1, value="ME do grupo"); ws.cell(row=r, column=6, value=P.ME).number_format = NUM; r += 1
ws.cell(row=r, column=1, value="Lucro não realizado nos estoques - base IR diferido 2025 (nota 20.9)"); ws.cell(row=r, column=6, value=1293806).number_format = NUM; r += 1
ws.cell(row=r, column=1, value="Lucro não realizado - base 2024"); ws.cell(row=r, column=6, value=1029612).number_format = NUM; r += 1
ws.cell(row=r, column=1, value="Variação"); ws.cell(row=r, column=6, value=f"=F{r-2}-F{r-1}").number_format = NUM; r += 1
ws.cell(row=r, column=1, value="Eliminação do LNR na equivalência 2025 (nota 9.3)"); ws.cell(row=r, column=6, value=-264234).number_format = NUM; r += 1
ws.cell(row=r, column=1, value="Diferença (variação da base x eliminação do exercício)"); ws.cell(row=r, column=6, value=f"=F{r-2}+F{r-1}").number_format = NUM
widths(ws, [62, 14, 14, 16, 12, 18, 52])

# ---------------------------------------------------------------- PT505
ws = wb.create_sheet("PT505")
title(ws, "PT 505 - Controle de circularização (MODELO - NENHUMA SOLICITAÇÃO FOI ENVIADA)", "NBC TA 505 itens 7-16 | Planejado por falta de acesso")
header(ws, 5, ["Ref.", "Parte a confirmar", "Saldo/assunto a confirmar (R$ mil)", "Fonte do saldo", "Tipo", "Data de envio",
               "Data de resposta", "Saldo confirmado", "Diferença", "Procedimento alternativo", "Status"])
lst = [("C01", "Banco Santander (Brasil) S.A. - empréstimo Res. 4131 e swap", "218.572 + swap passivo 22.729", "DF25 nota 15.1 p.40", "Positiva"),
       ("C02", "Vórtx DTVM (agente fiduciário) - 1ª emissão de debêntures", "312.704", "DF25 nota 15.1-15.2 p.40", "Positiva"),
       ("C03", "Banco Itaú BBA S.A. - confirmação de liquidação dos contratos de capital de giro", "0 (liquidados out/2025)", "DF25 nota 15.1 (a) p.40", "Positiva"),
       ("C04", "Instituições financeiras - caixa, aplicações compromissadas e letras financeiras", "398.644", "DF25 nota 5.2-5.3 p.29", "Positiva (nomes a obter)"),
       ("C05", "Adquirentes/operadoras de cartão - recebíveis", "986.130", "DF25 nota 6.2 p.29", "Positiva (nomes a obter)"),
       ("C06", "Instituições do convênio com fornecedores (Itaú/Santander)", "4.610", "DF25 nota 12.2 p.38", "Positiva"),
       ("C07", "Fornecedores nacionais e estrangeiros (amostra)", "117.227", "DF25 nota 12.1 p.38", "Positiva em branco"),
       ("C08", "Assessores jurídicos (cíveis, trabalhistas, tributários)", "Provisões 26.937; possíveis 382.523", "DF25 nota 17 p.41-45", "Carta a advogados"),
       ("C09", "Itaú Corretora (escriturador) - ações em tesouraria", "1.125.955 ações", "DF25 nota 18.2 p.45; FCA p.6", "Positiva"),
       ("C10", "Terceiros depositários / estoques em poder de terceiros (se houver)", "a identificar", "DF25 nota 7", "Positiva")]
for i, v in enumerate(lst, start=6):
    for j, x in enumerate(v, 1):
        ws.cell(row=i, column=j, value=x)
    ws.cell(row=i, column=6, value="não enviado"); ws.cell(row=i, column=7, value="-")
    ws.cell(row=i, column=10, value="Ver PT 505 (procedimentos alternativos)"); ws.cell(row=i, column=11, value="Planejado - não enviado (didático)")
widths(ws, [6, 60, 30, 26, 22, 14, 14, 14, 12, 34, 30])

fn = os.path.join(OUT, "PT_calculos_Vivara_241003384.xlsx")
wb.save(fn)
print("ok", fn)
